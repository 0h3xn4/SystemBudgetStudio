"""Provenance names every library that shapes an output, the packaged build can read their
versions, and every format dates itself with the provenance time (audit RC-5, RC-21)."""

from __future__ import annotations

import io
import re
import tomllib
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from budget_core.provenance import LIBRARIES, Provenance
from budget_core.reports.csvutil import defuse
from budget_core.reports.document import Column, ReportDocument, Section, Table
from budget_core.reports.docx import render_docx
from budget_core.reports.pdf import render_pdf
from budget_core.reports.xlsx import render_xlsx

ROOT = Path(__file__).resolve().parents[2]
PROV = Provenance(
    "System Budget Studio",
    "0.1.0",
    "3.13.0",
    (("numpy", "2.0"),),
    "Demo",
    "1",
    "s",
    "2026-03-04T05:06:07Z",
    "Jo Doe",
)


def dependency_names() -> set[str]:
    data = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    names = {re.split(r"[<>=!~\[ ]", d, maxsplit=1)[0] for d in data["project"]["dependencies"]}
    return {n.lower().replace("_", "-") for n in names}


def test_libraries_list_every_runtime_dependency_that_shapes_output() -> None:
    listed = {n.lower().replace("_", "-") for n in LIBRARIES}
    # the window toolkit draws the GUI only; every other runtime dependency shapes some output
    assert listed == {n for n in dependency_names() if not n.startswith("pyside6")}


def test_packaged_build_ships_the_metadata_of_every_listed_library() -> None:
    spec = (ROOT / "packaging" / "system_budget_studio.spec").read_text(encoding="utf-8")
    assert "copy_metadata" in spec and "LIBRARIES" in spec


def test_every_listed_library_has_a_version_in_this_environment() -> None:
    versions = dict(_provenance().libraries)
    assert set(versions) == set(LIBRARIES)
    assert "unknown" not in versions.values()


def _provenance() -> Provenance:
    from budget_core.io.project_loader import load_project
    from budget_core.provenance import make_provenance

    project = load_project(ROOT / "examples" / "microsat_250kg").project
    assert project is not None
    return make_provenance(project, user="u", generated_at=datetime(2026, 1, 1, tzinfo=UTC))


def document() -> ReportDocument:
    table = Table("T", (Column("a"),), (("x",),))
    return ReportDocument("Report", PROV, (Section("S", "S", tables=(table,)),))


def test_pdf_dates_are_the_provenance_time() -> None:
    meta = PdfReader(io.BytesIO(render_pdf(document()))).metadata
    assert meta is not None
    assert meta.creation_date == datetime(2026, 3, 4, 5, 6, 7, tzinfo=UTC)
    assert meta.modification_date == datetime(2026, 3, 4, 5, 6, 7, tzinfo=UTC)
    assert render_pdf(document()) == render_pdf(document())


def test_docx_app_properties_name_this_tool() -> None:
    data = render_docx(document())
    app = zipfile.ZipFile(io.BytesIO(data)).read("docProps/app.xml").decode("utf-8")
    assert "Macintosh" not in app and "System Budget Studio" in app


def test_docx_core_properties_are_truncated_instead_of_failing() -> None:
    long = "x" * 400
    prov = Provenance("t", "0", "3", (), long, "1", "s", "2026-01-01T00:00:00Z", long)
    doc = ReportDocument(long, prov, (Section("S", "S"),))
    from docx import Document

    props = Document(io.BytesIO(render_docx(doc))).core_properties
    assert len(props.author) <= 255 and len(props.title) <= 255 and len(props.subject) <= 255


def test_xlsx_creator_is_the_user_like_the_other_formats() -> None:
    wb = load_workbook(io.BytesIO(render_xlsx(document())))
    assert wb.properties.creator == "Jo Doe" and wb.properties.lastModifiedBy == (
        "System Budget Studio"
    )


@pytest.mark.parametrize("text", ["-Jo", "+Anna-Lisa Smith", "-Müller", "plain", "1+1 plan"])
def test_names_that_cannot_run_are_not_rewritten(text: str) -> None:
    assert defuse(text) == text


@pytest.mark.parametrize("text", ["=1+1", "@SUM(A1)", "+cmd|x", "-1+cmd|' /C calc'!A0", "\tx"])
def test_dangerous_leading_text_gets_one_apostrophe(text: str) -> None:
    assert defuse(text) == "'" + text


def test_scenario_json_provenance_names_python_and_libraries() -> None:
    import json

    from budget_core.scenario.export import environment_json
    from tests.power_helpers import synthetic_env, synthetic_run

    run = synthetic_run(synthetic_env(600.0, 60.0))
    data = json.loads(environment_json(run, PROV))["provenance"]
    assert data["python_version"] == "3.13.0" and data["libraries"] == [["numpy", "2.0"]]
