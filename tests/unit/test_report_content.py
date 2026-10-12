"""What the reports contain: load-time problems in JSON, Path and Hint columns, the INCOMPLETE
banner on every sheet and page, equation lists, link banners, sheet names and renderer limits
(audit RC-4, RC-11, RC-13, RC-14, RC-19, RC-25, TL-3)."""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from budget_core.io.project_loader import load_project
from budget_core.link.evaluate import link_pass_series, link_static_budget
from budget_core.model import Ebn0Table
from budget_core.power.static_budget import static_power_budget
from budget_core.power.time_domain import time_domain_budget
from budget_core.problems import Problem, Severity
from budget_core.provenance import Provenance, make_provenance
from budget_core.reports.common import BANNER, sheet_name
from budget_core.reports.document import Column, ReportDocument, Section, Table
from budget_core.reports.docx import render_docx
from budget_core.reports.pdf import render_pdf
from budget_core.reports.run import (
    BudgetOutput,
    link_output,
    power_output,
    timeline_output,
    write_outputs,
)
from budget_core.reports.xlsx import render_xlsx
from tests.link_helpers import downlink, link_project, pass_run
from tests.power_helpers import synthetic_env, synthetic_run, td_project

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
WHEN = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)
PROV = Provenance("t", "0", "3", (), "P", "1", "s", "2026-01-01T00:00:00Z", "u")
LOAD = Problem(
    Severity.WARNING,
    "FILE_IGNORED",
    "A file was ignored.",
    file="units/x.yaml",
    path="modes[0]",
    hint="Rename the file.",
)


def cubesat_output(load_problems: list[Problem]) -> BudgetOutput:
    project = load_project(EXAMPLES / "cubesat_3u").project
    assert project is not None
    prov = make_provenance(project, user="u", generated_at=WHEN)
    return power_output(project, static_power_budget(project), prov, load_problems)


# ---- RC-4 -----------------------------------------------------------------------------------
def test_json_problems_include_the_load_time_problems(tmp_path: Path) -> None:
    out = cubesat_output([LOAD])
    write_outputs([out], tmp_path, {"json"})
    data = json.loads((tmp_path / "power_static.json").read_text(encoding="utf-8"))
    codes = [p["code"] for p in data["problems"]]
    assert "FILE_IGNORED" in codes
    table = next(
        t for s in out.document.sections for t in s.tables if t.title == "Problems and open items"
    )
    assert len(codes) == len(table.rows)  # the JSON lists what the XLSX, PDF and DOCX list


def test_timeline_and_link_json_include_the_load_time_problems() -> None:
    project = td_project()
    result = time_domain_budget(project, synthetic_run(synthetic_env(600.0, 60.0)))
    prov = make_provenance(project, user="u", generated_at=WHEN)
    out = timeline_output(project, result, prov, [LOAD], plots=False)
    assert "FILE_IGNORED" in [p["code"] for p in json.loads(out.json_text or "")["problems"]]
    lp = link_project()
    static = link_static_budget(lp)
    out = link_output(lp, static, prov, [LOAD], plots=False)
    assert "FILE_IGNORED" in [p["code"] for p in json.loads(out.json_text or "")["problems"]]


# ---- RC-11 ----------------------------------------------------------------------------------
def test_problems_table_has_path_and_hint_columns() -> None:
    out = cubesat_output([LOAD])
    table = next(
        t for s in out.document.sections for t in s.tables if t.title == "Problems and open items"
    )
    headers = [c.header for c in table.columns]
    assert "Path" in headers and "Hint" in headers
    row = next(r for r in table.rows if r[1] == "FILE_IGNORED")
    assert row[headers.index("Path")] == "modes[0]"
    assert row[headers.index("Hint")] == "Rename the file."


# ---- RC-13 ----------------------------------------------------------------------------------
def banner_document() -> ReportDocument:
    sections = tuple(
        Section(f"S{i}", f"S{i}", tables=(Table("T", (Column("a"),), (("x",),) * 3),))
        for i in range(3)
    )
    return ReportDocument("R", PROV, sections, BANNER)


def test_xlsx_banner_is_on_every_sheet() -> None:
    wb = load_workbook(io.BytesIO(render_xlsx(banner_document())))
    for ws in wb:
        assert str(ws["A1"].value).startswith("INCOMPLETE") or str(ws["A2"].value).startswith(
            "INCOMPLETE"
        ), ws.title


def test_pdf_banner_is_on_every_page() -> None:
    table = Table("T", (Column("a"),), (("x",),) * 120)  # several pages
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),), BANNER)
    pages = PdfReader(io.BytesIO(render_pdf(doc))).pages
    assert len(pages) > 2
    assert all("INCOMPLETE" in p.extract_text() for p in pages)


def test_docx_banner_is_in_the_header_of_every_page() -> None:
    document = Document(io.BytesIO(render_docx(banner_document())))
    header = "\n".join(p.text for p in document.sections[0].header.paragraphs)
    assert "INCOMPLETE" in header


def test_complete_reports_have_no_banner_anywhere() -> None:
    doc = replace(banner_document(), banner="")
    assert all(
        "INCOMPLETE" not in p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(doc))).pages
    )
    wb = load_workbook(io.BytesIO(render_xlsx(doc)))
    assert all("INCOMPLETE" not in str(c.value) for ws in wb for r in ws.iter_rows() for c in r)


def test_json_carries_the_incomplete_flag(tmp_path: Path) -> None:
    out = cubesat_output([])
    write_outputs([out], tmp_path, {"json"})
    data = json.loads((tmp_path / "power_static.json").read_text(encoding="utf-8"))
    assert data["incomplete"] is True


# ---- RC-14 ----------------------------------------------------------------------------------
def equation_ids(doc: ReportDocument) -> set[str]:
    table = next(t for s in doc.sections for t in s.tables if t.title == "Equations")
    return {str(r[0]) for r in table.rows}


def test_timeline_report_lists_power_and_environment_equations() -> None:
    project = td_project()
    result = time_domain_budget(project, synthetic_run(synthetic_env(600.0, 60.0)))
    prov = make_provenance(project, user="u", generated_at=WHEN)
    ids = equation_ids(timeline_output(project, result, prov, plots=False).document)
    assert {"TDP-ARRAY", "PWR-SOURCE", "PWR-MARGIN", "ENV-SUN"} <= ids
    assert ("ENV-SHADOW-CYL" in ids) != ("ENV-SHADOW-CON" in ids)  # the model in use only


def test_link_pass_report_lists_environment_equations_the_static_one_does_not() -> None:
    project = link_project()
    prov = make_provenance(project, user="u", generated_at=WHEN)
    static = link_static_budget(project)
    passes = link_pass_series(project, pass_run())
    with_passes = equation_ids(
        link_output(project, static, prov, passes=passes, plots=False).document
    )
    without = equation_ids(link_output(project, static, prov, plots=False).document)
    assert {"ENV-TOPO", "ENV-GEODETIC"} <= with_passes and not {"ENV-TOPO"} & without
    assert {"LNK-FSPL"} <= without


# ---- TL-3 -----------------------------------------------------------------------------------
def test_link_report_has_a_banner_when_pass_results_are_n_a_for_a_missing_table_entry() -> None:
    project = link_project(links={"dl": downlink(static_points=[])})
    project = replace(project, config=replace(project.config, ebn0_table=Ebn0Table(entries=[])))
    prov = make_provenance(project, user="u", generated_at=WHEN)
    static = link_static_budget(project)
    passes = link_pass_series(project, pass_run())
    assert passes.series[0].margin_db is None
    out = link_output(project, static, prov, passes=passes, plots=False)
    assert "INCOMPLETE" in out.document.banner and out.incomplete


def test_link_report_without_problems_has_no_banner() -> None:
    project = link_project()
    prov = make_provenance(project, user="u", generated_at=WHEN)
    passes = link_pass_series(project, pass_run())
    out = link_output(project, link_static_budget(project), prov, passes=passes, plots=False)
    assert out.document.banner == ""


# ---- RC-25 ----------------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Ann's", "Ann's"),
        ("'quoted'", "quoted"),
        ("a'", "a"),
        ("", "Sheet"),
        ("History", "History_"),
    ],
)
def test_sheet_name_rules(raw: str, expected: str) -> None:
    assert sheet_name(raw, set()) == expected


def test_sheet_names_are_unique_ignoring_case() -> None:
    used: set[str] = set()
    names = [sheet_name(n, used) for n in ("Plots", "plots", "PLOTS", "x" * 40, "x" * 40)]
    assert len({n.lower() for n in names}) == 5 and all(len(n) <= 31 for n in names)


def test_section_sheet_names_are_the_real_sheet_names() -> None:
    used: set[str] = set()
    sections = tuple(
        Section(raw, sheet_name(raw, used), tables=(Table("T", (Column("a"),), (("x",),)),))
        for raw in ("Plots", "plots", "'Odd'", "a/b:c", "x" * 50, "x" * 50, "History")
    )
    wb = load_workbook(io.BytesIO(render_xlsx(ReportDocument("R", PROV, sections))))
    assert wb.sheetnames == [s.sheet_name for s in sections]


def test_a_very_long_cell_is_shortened_with_a_note_in_the_spreadsheet() -> None:
    table = Table("T", (Column("a"),), (("y" * 40000,),))
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    sheet = load_workbook(io.BytesIO(render_xlsx(doc)))["S"]
    texts = [str(c.value) for r in sheet.iter_rows() for c in r if c.value]
    assert max(len(t) for t in texts) <= 32767
    assert any("shortened" in t for t in texts)


# ---- RC-19 ----------------------------------------------------------------------------------
def pdf_text(doc: ReportDocument) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(doc))).pages)


def test_pdf_survives_a_very_long_cell_and_says_so() -> None:
    table = Table("T", (Column("a"), Column("b")), (("z" * 10000, "short"),))
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    text = pdf_text(doc)
    assert "shortened" in text and "short" in text


def test_pdf_survives_many_long_cells_in_one_row() -> None:
    table = Table(
        "T", tuple(Column(f"c{i}") for i in range(8)), (tuple("w " * 3000 for _ in range(8)),)
    )
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    assert "shortened" in pdf_text(doc)


def test_pdf_header_words_are_not_split() -> None:
    columns = tuple(Column(h) for h in ("Unit", "Subsystem", "Bus", "Power mode"))
    table = Table("T", columns, (("a", "b", "c", "d"),))
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    assert "Subsystem" in pdf_text(doc)


def test_pdf_shows_unsupported_characters_as_question_marks_with_a_note() -> None:
    table = Table("T", (Column("a"),), (("日本語 ok",),))
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    text = pdf_text(doc)
    assert "??? ok" in text and "not available in the PDF font" in text


def test_pdf_long_footer_does_not_run_into_the_page_number() -> None:
    prov = replace(PROV, project_name="N" * 300, user="U" * 200)
    doc = ReportDocument("R", prov, (Section("S", "S"),))
    text = pdf_text(doc)
    assert "Page 1" in text
