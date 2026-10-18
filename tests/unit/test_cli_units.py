"""`budget units-export` and `budget units-import`: a template, a dry run, then --apply."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pytest

from budget_cli.main import main
from budget_core.io.project_loader import load_project

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"


@pytest.fixture
def root(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES / "microsat_250kg_eps", tmp_path / "p")
    return tmp_path / "p"


def edited_table(root: Path, tmp_path: Path, new_avg: str = "30.0 W") -> Path:
    assert main(["units-export", str(root), "--out", str(tmp_path / "t.csv")]) == 0
    path = tmp_path / "t.csv"
    rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    header = rows[0]
    for r in rows[1:]:
        if r[0] == "radio" and r[header.index("power_mode")] == "downlink":
            r[header.index("avg_power_w")] = new_avg
    with path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(rows)
    return path


def test_export_writes_a_csv_template_next_to_the_results_by_default(
    root: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["units-export", str(root)]) == 0
    target = root / "results" / "unit_power_modes.csv"
    assert target.is_file() and str(target) in capsys.readouterr().out
    assert target.read_bytes().startswith(b"\xef\xbb\xbf")  # a BOM, so Excel reads the UTF-8


def test_export_can_write_an_xlsx_workbook(root: Path, tmp_path: Path) -> None:
    out = tmp_path / "t.xlsx"
    assert main(["units-export", str(root), "--out", str(out)]) == 0
    from openpyxl import load_workbook

    assert load_workbook(out).sheetnames[0] == "Unit power modes"


def test_export_refuses_a_file_type_it_does_not_write(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["units-export", str(root), "--out", str(tmp_path / "t.ods")]) == 2
    assert ".csv" in capsys.readouterr().err


def test_import_shows_the_changes_and_writes_nothing_without_apply(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    table = edited_table(root, tmp_path)
    before = (root / "units" / "radio.yaml").read_bytes()
    capsys.readouterr()
    assert main(["units-import", str(root), str(table)]) == 0
    out = capsys.readouterr().out
    assert "radio" in out and "downlink" in out and "avg_power_w" in out
    assert "28" in out and "30" in out
    assert "Nothing was written" in out and "--apply" in out
    assert (root / "units" / "radio.yaml").read_bytes() == before


def test_apply_writes_the_unit_file_and_the_project_still_loads(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    table = edited_table(root, tmp_path)
    capsys.readouterr()
    assert main(["units-import", str(root), str(table), "--apply"]) == 0
    assert "units/radio.yaml" in capsys.readouterr().out.replace("\\", "/")
    project = load_project(root).project
    assert project is not None and project.units["radio"].modes[5].avg_power_w == pytest.approx(
        30.0
    )


def test_an_untouched_template_reports_no_change(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["units-export", str(root), "--out", str(tmp_path / "t.csv")]) == 0
    capsys.readouterr()
    assert main(["units-import", str(root), str(tmp_path / "t.csv"), "--apply"]) == 0
    assert "No change" in capsys.readouterr().out


def test_a_problem_is_listed_exits_1_and_writes_nothing(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    table = edited_table(root, tmp_path, new_avg="a lot")
    before = (root / "units" / "radio.yaml").read_bytes()
    capsys.readouterr()
    assert main(["units-import", str(root), str(table), "--apply"]) == 1
    out = capsys.readouterr().err
    assert "IMPORT_VALUE_INVALID" in out and "a lot" not in out
    assert "nothing was written" in out.lower()
    assert (root / "units" / "radio.yaml").read_bytes() == before


def test_a_missing_table_is_a_plain_message(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["units-import", str(root), str(tmp_path / "nope.csv")]) == 1
    out = capsys.readouterr().err
    assert "IMPORT_FILE_UNREADABLE" in out and "Traceback" not in out


def test_a_project_with_errors_is_not_touched(
    root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    table = edited_table(root, tmp_path)
    (root / "spacecraft.yaml").unlink()
    capsys.readouterr()
    assert main(["units-import", str(root), str(table), "--apply"]) == 1
    assert "nothing was written" in capsys.readouterr().err.lower()


def test_apply_works_with_a_relative_project_path(
    root: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    table = edited_table(root, tmp_path)
    monkeypatch.chdir(tmp_path)
    capsys.readouterr()
    assert main(["units-import", "p", "t.csv", "--apply"]) == 0
    assert "Wrote units/radio.yaml" in capsys.readouterr().out
    assert table.exists()
