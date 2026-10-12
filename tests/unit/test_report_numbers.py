"""How report numbers are shown: round half up on the decimal value, no negative zero, no loss in
the configuration-numbers table, the same text in every format (audit RC-8, RC-9, TL-16, RC-25)."""

from __future__ import annotations

import io
import json
import math

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from budget_core.assumptions import Assumption
from budget_core.provenance import Provenance
from budget_core.reports.common import closing_sections
from budget_core.reports.document import NA, Column, ReportDocument, Section, Table, format_cell
from budget_core.reports.docx import render_docx
from budget_core.reports.export import result_json
from budget_core.reports.pdf import render_pdf
from budget_core.reports.xlsx import render_xlsx

PROV = Provenance("t", "0", "3", (), "P", "1", "s", "2026-01-01T00:00:00Z", "u")


@pytest.mark.parametrize(
    ("value", "decimals", "expected"),
    [
        (0.125, 2, "0.13"),  # the binary value is exact; Excel shows 0.13
        (2.675, 2, "2.68"),  # the binary value is 2.67499999...; Excel shows 2.68
        (-0.125, 2, "-0.13"),
        (0.135, 2, "0.14"),
        (1.005, 2, "1.01"),
        (0.5, 0, "1"),
        (-0.5, 0, "-1"),
        (-0.0, 2, "0.00"),
        (-0.001, 2, "0.00"),
        (-0.0049999, 2, "0.00"),
        (1e22, 1, "10000000000000000000000.0"),
        (123456.789, 1, "123456.8"),
    ],
)
def test_numbers_round_half_up_on_the_decimal_value(
    value: float, decimals: int, expected: str
) -> None:
    assert format_cell(value, Column("x", "number", decimals)) == expected


def test_percent_rounds_half_up_and_has_no_negative_zero() -> None:
    column = Column("x", "percent", 3)  # one decimal in percent
    assert format_cell(0.12345, column) == "12.3 %"
    assert format_cell(0.1235, column) == "12.4 %"  # 12.35 exactly
    assert format_cell(-0.00001, column) == "0.0 %"


def test_non_finite_numbers_are_n_a() -> None:
    for value in (math.nan, math.inf, -math.inf):
        assert format_cell(value, Column("x", "number", 2)) == NA


def test_general_columns_lose_nothing() -> None:
    column = Column("x", "number", -1)
    assert format_cell(-0.00041234, column) == "-0.00041234"
    assert format_cell(1e-9, column) == "1e-09"
    assert format_cell(0.0, column) == "0"
    assert format_cell(-0.0, column) == "0"
    assert format_cell(1234.5, column) == "1234.5"


def one_cell_document(value: float, decimals: int = 2) -> ReportDocument:
    table = Table("T", (Column("Value", "number", decimals),), ((value,),))
    return ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))


def test_pdf_docx_and_xlsx_show_the_same_rounding() -> None:
    doc = one_cell_document(0.125)
    pdf = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(render_pdf(doc))).pages)
    assert "0.13" in pdf and "0.12" not in pdf
    docx = Document(io.BytesIO(render_docx(doc)))
    assert docx.tables[0].rows[1].cells[0].text == "0.13"
    sheet = load_workbook(io.BytesIO(render_xlsx(doc)))["S"]
    cell = sheet["A5"]
    assert cell.value == 0.125 and cell.number_format == "0.00"  # Excel rounds half up too


def test_xlsx_stores_no_negative_zero_and_n_a_for_non_finite() -> None:
    table = Table(
        "T",
        (Column("a", "number", 2), Column("b", "number", 2), Column("c", "number", 2)),
        ((-0.0, math.nan, math.inf),),
    )
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    sheet = load_workbook(io.BytesIO(render_xlsx(doc)))["S"]
    values = [c.value for c in sheet[5]]
    assert values[0] == 0 and math.copysign(1.0, values[0]) == 1.0
    assert values[1:] == [NA, NA]


def test_xlsx_percent_format_matches_the_displayed_decimals() -> None:
    table = Table("T", (Column("p", "percent", 3), Column("q", "percent", 2)), ((0.1234, 0.5),))
    doc = ReportDocument("R", PROV, (Section("S", "S", tables=(table,)),))
    sheet = load_workbook(io.BytesIO(render_xlsx(doc)))["S"]
    assert sheet["A5"].number_format == "0.0%" and sheet["B5"].number_format == "0%"
    assert format_cell(0.1234, Column("p", "percent", 3)) == "12.3 %"
    assert format_cell(0.5, Column("q", "percent", 2)) == "50 %"


def test_configuration_numbers_table_shows_values_without_loss() -> None:
    assumption = Assumption("k", -0.00041234, "1/K", "ref", "f.yaml", "k", False)
    tiny = Assumption("e", 3e-9, "W", "ref", "f.yaml", "e", False)
    sections = closing_sections(set(), [assumption, tiny], "PWR-", [], PROV)
    table = sections[0].tables[0]
    shown = [format_cell(row[1], table.columns[1]) for row in table.rows]
    assert shown == ["-0.00041234", "3e-09"]
    docx = Document(io.BytesIO(render_docx(ReportDocument("R", PROV, tuple(sections)))))
    cells = [c.text for row in docx.tables[0].rows for c in row.cells]
    assert "-0.00041234" in cells and "3e-09" in cells


def test_json_has_no_negative_zero() -> None:
    from dataclasses import dataclass

    @dataclass
    class Result:
        value: float
        problems: list[object]

    text = result_json(Result(-0.0, []), PROV, incomplete=False)
    assert json.loads(text)["value"] == 0 and "-0.0" not in text


def test_csv_cells_have_no_negative_zero() -> None:
    from budget_core.reports.export import _cell

    assert _cell(-0.0) == "0.0"
