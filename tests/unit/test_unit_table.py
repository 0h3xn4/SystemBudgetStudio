"""Unit power data in a table: export a template, fill it in a spreadsheet, import it back.

An import first shows what it would change; nothing is written unless it is applied, and then
all of it or none of it. Empty cells mean "unchanged", never zero. Messages never repeat values
read from a file.
"""

from __future__ import annotations

import csv
import io
import shutil
from dataclasses import replace
from pathlib import Path

import pytest

from budget_core.io.project_loader import load_project
from budget_core.io.unit_table import (
    HEADER,
    apply_plan,
    plan_import,
    read_table,
    table_csv,
    table_xlsx,
)
from budget_core.model import Project
from budget_core.problems import Severity

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"


@pytest.fixture
def root(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES / "microsat_250kg_eps", tmp_path / "p")
    return tmp_path / "p"


def project_of(root: Path) -> Project:
    loaded = load_project(root)
    assert loaded.project is not None
    return loaded.project


def rows_of(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def write_table(path: Path, rows: list[list[object]], delimiter: str = ",") -> Path:
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=delimiter, lineterminator="\n").writerows(rows)
    path.write_text("﻿" + buffer.getvalue(), encoding="utf-8")
    return path


def import_rows(root: Path, tmp_path: Path, *data_rows: list[object], header=None):  # type: ignore[no-untyped-def]
    head = list(header or HEADER)
    table = read_table(write_table(tmp_path / "in.csv", [head, *data_rows]))
    return plan_import(project_of(root), table, "in.csv")


def row(unit: str, mode: str, **values: object) -> list[object]:
    return [
        values.get(c, "")
        if c not in ("unit_id", "power_mode")
        else {"unit_id": unit, "power_mode": mode}[c]
        for c in HEADER
    ]


# ---- export -----------------------------------------------------------------------------------


def test_the_template_has_one_row_per_power_mode_in_file_order() -> None:
    project = load_project(EXAMPLES / "microsat_250kg_eps").project
    assert project is not None
    rows = rows_of(table_csv(project))
    assert tuple(rows[0]) == HEADER
    expected = sum(len(u.modes) for u in project.units.values())
    assert len(rows) - 1 == expected
    radio = [r for r in rows[1:] if r[0] == "radio"]
    assert [r[HEADER.index("power_mode")] for r in radio] == [
        m.name for m in project.units["radio"].modes
    ]
    downlink = next(r for r in radio if r[HEADER.index("power_mode")] == "downlink")
    assert float(downlink[HEADER.index("avg_power_w")]) == pytest.approx(28.0)
    assert float(downlink[HEADER.index("duty_cycle_ratio")]) == pytest.approx(1.0)


def test_the_header_names_carry_their_units() -> None:
    assert {"avg_power_w", "peak_power_w", "duty_cycle_ratio", "heat_dissipation_ratio"} <= set(
        HEADER
    )
    assert {"min_duration_s", "max_duration_s"} <= set(HEADER)


def test_text_that_a_spreadsheet_would_run_as_a_formula_is_defused(root: Path) -> None:
    project = project_of(root)
    unit = project.units["radio"]
    renamed = unit.model_copy(update={"name": "=cmd|' /C calc'!A0"})
    project = replace(project, units={**project.units, "radio": renamed})
    csv_text = table_csv(project)
    assert "'=cmd" in csv_text and ",=cmd" not in csv_text
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(table_xlsx(project))).active
    names = [c.value for c in sheet["B"] if c.value]
    assert any(str(n).startswith("'=") for n in names)
    assert all(c.data_type != "f" for r in sheet.iter_rows() for c in r)


# ---- import: what it plans --------------------------------------------------------------------


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_importing_the_untouched_template_changes_nothing(
    root: Path, tmp_path: Path, kind: str
) -> None:
    project = project_of(root)
    if kind == "csv":
        path = tmp_path / "t.csv"
        path.write_text(table_csv(project), encoding="utf-8")
    else:
        path = tmp_path / "t.xlsx"
        path.write_bytes(table_xlsx(project))
    plan = plan_import(project, read_table(path), path.name)
    assert plan.ok and plan.changes == () and not plan.problems
    assert plan.unchanged == sum(len(u.modes) for u in project.units.values())


def test_a_changed_value_is_planned_with_old_and_new(root: Path, tmp_path: Path) -> None:
    plan = import_rows(root, tmp_path, row("radio", "downlink", avg_power_w="30.0 W"))
    assert plan.ok
    (change,) = plan.changes
    assert (change.unit_id, change.power_mode, change.field) == ("radio", "downlink", "avg_power_w")
    assert change.old == pytest.approx(28.0) and change.new == pytest.approx(30.0)
    assert change.kind == "changed"


def test_units_are_converted_like_in_the_yaml_files(root: Path, tmp_path: Path) -> None:
    plan = import_rows(
        root,
        tmp_path,
        row("radio", "downlink", avg_power_w="800 mW", peak_power_w=3.5, duty_cycle_ratio="40 %"),
    )
    assert plan.ok
    values = {c.field: c.new for c in plan.changes}
    assert values == pytest.approx(
        {"avg_power_w": 0.8, "peak_power_w": 3.5, "duty_cycle_ratio": 0.4}
    )


def test_an_empty_cell_means_unchanged_not_zero(root: Path, tmp_path: Path) -> None:
    plan = import_rows(root, tmp_path, row("radio", "downlink", peak_power_w=45.0))
    assert [c.field for c in plan.changes] == ["peak_power_w"]


def test_a_value_equal_to_the_current_one_is_not_a_change(root: Path, tmp_path: Path) -> None:
    plan = import_rows(
        root, tmp_path, row("radio", "downlink", avg_power_w=28.0, peak_power_w="40000 mW")
    )
    assert plan.ok and plan.changes == () and plan.unchanged == 1


def test_a_new_power_mode_needs_average_and_peak(root: Path, tmp_path: Path) -> None:
    good = import_rows(root, tmp_path, row("radio", "burst", avg_power_w=3.0, peak_power_w=5.0))
    assert good.ok
    assert {c.kind for c in good.changes} == {"added"}
    assert {c.field for c in good.changes} >= {"avg_power_w", "peak_power_w"}
    bad = import_rows(root, tmp_path, row("radio", "burst", avg_power_w=3.0))
    assert not bad.ok
    (problem,) = [p for p in bad.problems if p.severity is Severity.ERROR]
    assert problem.code == "IMPORT_FIELD_MISSING" and problem.path == "peak_power_w"
    assert problem.file == "in.csv" and problem.line == 2


def test_the_semicolon_delimiter_and_a_byte_order_mark_are_accepted(
    root: Path, tmp_path: Path
) -> None:
    path = write_table(
        tmp_path / "semi.csv",
        [list(HEADER), row("radio", "downlink", avg_power_w="30.0 W")],
        delimiter=";",
    )
    plan = plan_import(project_of(root), read_table(path), "semi.csv")
    assert plan.ok and [c.field for c in plan.changes] == ["avg_power_w"]


# ---- import: refusals -------------------------------------------------------------------------


def test_an_unknown_unit_is_an_error_and_nothing_is_planned_for_it(
    root: Path, tmp_path: Path
) -> None:
    plan = import_rows(
        root,
        tmp_path,
        row("radio", "downlink", avg_power_w="30.0 W"),
        row("ghost", "on", avg_power_w=1.0, peak_power_w=2.0),
    )
    assert not plan.ok
    (problem,) = [p for p in plan.problems if p.severity is Severity.ERROR]
    assert problem.code == "IMPORT_UNKNOWN_UNIT" and problem.line == 3 and problem.path == "unit_id"
    assert "ghost" not in problem.message  # values from the file are never repeated


def test_a_repeated_unit_and_mode_is_an_error(root: Path, tmp_path: Path) -> None:
    plan = import_rows(
        root,
        tmp_path,
        row("radio", "downlink", avg_power_w=30.0),
        row("radio", "downlink", avg_power_w=31.0),
    )
    assert any(p.code == "IMPORT_DUPLICATE_ROW" and p.line == 3 for p in plan.problems)


def test_a_value_the_model_rejects_is_located_without_echoing_it(
    root: Path, tmp_path: Path
) -> None:
    plan = import_rows(
        root, tmp_path, row("radio", "downlink", avg_power_w=9.123456, peak_power_w=1.0)
    )
    assert not plan.ok
    (problem,) = [p for p in plan.problems if p.severity is Severity.ERROR]
    assert problem.code == "IMPORT_VALUE_INVALID" and problem.line == 2
    assert problem.path == "peak_power_w"
    assert "9.123456" not in problem.message + problem.hint


@pytest.mark.parametrize("cell", ["abc", "3 dBW", "1,5", "nan", "-1 W"])
def test_text_that_is_not_a_number_with_a_matching_unit_is_refused(
    root: Path, tmp_path: Path, cell: str
) -> None:
    plan = import_rows(root, tmp_path, row("radio", "downlink", avg_power_w=cell))
    assert not plan.ok
    assert all(
        p.code == "IMPORT_VALUE_INVALID" for p in plan.problems if p.severity is Severity.ERROR
    )
    assert cell not in " ".join(p.message + p.hint for p in plan.problems)


def test_a_table_without_the_key_columns_is_refused(root: Path, tmp_path: Path) -> None:
    path = write_table(tmp_path / "x.csv", [["avg_power_w"], ["1.0"]])
    plan = plan_import(project_of(root), read_table(path), "x.csv")
    assert not plan.ok
    codes = {p.code for p in plan.problems}
    assert "IMPORT_COLUMN_MISSING" in codes


def test_unknown_columns_are_ignored_with_a_warning_that_names_the_position(
    root: Path, tmp_path: Path
) -> None:
    header = [*HEADER, "supplier comment"]
    plan = import_rows(
        root, tmp_path, [*row("radio", "downlink", avg_power_w=30.0), "check!"], header=header
    )
    assert plan.ok and [c.field for c in plan.changes] == ["avg_power_w"]
    (warning,) = [p for p in plan.problems if p.severity is Severity.WARNING]
    assert warning.code == "IMPORT_COLUMN_IGNORED" and "supplier" not in warning.message


def test_a_file_that_is_too_large_or_not_a_table_is_refused(tmp_path: Path) -> None:
    big = tmp_path / "big.csv"
    big.write_text("unit_id,power_mode\n" + "x,y\n" * 2_000_000, encoding="utf-8")
    assert any(p.code == "IMPORT_FILE_TOO_LARGE" for p in read_table(big).problems)
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a zip file")
    assert any(p.code == "IMPORT_FILE_UNREADABLE" for p in read_table(junk).problems)
    assert any(
        p.code == "IMPORT_FILE_UNREADABLE" for p in read_table(tmp_path / "missing.csv").problems
    )
    odd = tmp_path / "t.txt"
    odd.write_text("a", encoding="utf-8")
    assert any(p.code == "IMPORT_FILE_UNREADABLE" for p in read_table(odd).problems)


def test_a_formula_in_a_spreadsheet_is_never_evaluated(root: Path, tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(list(HEADER))
    ws.append(["radio", "", "", "downlink", "=1+1", "", "", "", "", ""][: len(HEADER)])
    path = tmp_path / "f.xlsx"
    wb.save(path)
    plan = plan_import(project_of(root), read_table(path), "f.xlsx")
    assert plan.changes == ()  # no cached value: the cell counts as empty, not as 2


# ---- apply ------------------------------------------------------------------------------------


def test_planning_writes_nothing(root: Path, tmp_path: Path) -> None:
    before = (root / "units" / "radio.yaml").read_bytes()
    import_rows(root, tmp_path, row("radio", "downlink", avg_power_w="30.0 W"))
    assert (root / "units" / "radio.yaml").read_bytes() == before


def test_applying_changes_only_the_named_values_and_keeps_comments(
    root: Path, tmp_path: Path
) -> None:
    path = root / "units" / "radio.yaml"
    text = path.read_text(encoding="utf-8")
    path.write_text(
        text.replace(
            "    avg_power_w: 28.0\n", "    # supplier data sheet rev B\n    avg_power_w: 28.0\n"
        ),
        encoding="utf-8",
    )
    original = path.read_text(encoding="utf-8")
    plan = import_rows(root, tmp_path, row("radio", "downlink", avg_power_w="30.0 W"))
    written = apply_plan(root, plan)
    assert written == [path]
    after = path.read_text(encoding="utf-8")
    assert "# supplier data sheet rev B" in after
    changed = [
        (a, b) for a, b in zip(original.splitlines(), after.splitlines(), strict=True) if a != b
    ]
    assert changed == [("    avg_power_w: 28.0", "    avg_power_w: 30.0")]
    assert project_of(root).units["radio"].modes[5].avg_power_w == pytest.approx(30.0)
    assert (root / "units" / "obc.yaml").read_bytes() == (
        EXAMPLES / "microsat_250kg_eps" / "units" / "obc.yaml"
    ).read_bytes()


def test_applying_a_new_mode_appends_it_and_the_project_still_loads(
    root: Path, tmp_path: Path
) -> None:
    plan = import_rows(
        root,
        tmp_path,
        row("radio", "burst", avg_power_w=3.0, peak_power_w=5.0, heat_dissipation_ratio=0.6),
    )
    apply_plan(root, plan)
    loaded = load_project(root)
    names = [m.name for m in project_of(root).units["radio"].modes]
    assert names[-1] == "burst" and names[:-1] == [
        m.name
        for m in load_project(EXAMPLES / "microsat_250kg_eps").project.units["radio"].modes  # type: ignore[union-attr]
    ]
    burst = project_of(root).units["radio"].modes[-1]
    assert (burst.avg_power_w, burst.peak_power_w, burst.heat_dissipation_ratio) == (3.0, 5.0, 0.6)
    assert not [p for p in loaded.problems if p.severity is Severity.ERROR]


def test_a_plan_with_an_error_writes_nothing_at_all(root: Path, tmp_path: Path) -> None:
    before = {p: p.read_bytes() for p in (root / "units").glob("*.yaml")}
    plan = import_rows(
        root,
        tmp_path,
        row("radio", "downlink", avg_power_w="30.0 W"),
        row("ghost", "on", avg_power_w=1.0, peak_power_w=2.0),
    )
    assert apply_plan(root, plan) == []
    assert {p: p.read_bytes() for p in (root / "units").glob("*.yaml")} == before


def test_applying_twice_is_the_same_as_once(root: Path, tmp_path: Path) -> None:
    plan = import_rows(root, tmp_path, row("radio", "downlink", avg_power_w="30.0 W"))
    apply_plan(root, plan)
    once = (root / "units" / "radio.yaml").read_bytes()
    again = import_rows(root, tmp_path, row("radio", "downlink", avg_power_w="30.0 W"))
    assert again.changes == ()
    apply_plan(root, again)
    assert (root / "units" / "radio.yaml").read_bytes() == once
