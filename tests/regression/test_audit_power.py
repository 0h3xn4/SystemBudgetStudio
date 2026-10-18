"""Pre-release audit fixes in the time-domain power budget (PWR-1 to PWR-12).

Numbers are invented and worked out by hand (see `tests/power_helpers.py`: 100 W array at cosine 1,
100 Wh battery, mode a 40 W, mode b 100 W).
"""

from __future__ import annotations

import json
import math
import shutil
from dataclasses import replace
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from budget_cli.main import main
from budget_core.environment.data import EnvironmentData
from budget_core.environment.intervals import Interval
from budget_core.model import ArrayFace, Attitude, Battery, ScenarioRule
from budget_core.power.orbits import orbit_windows, orbit_windows_with_reason
from budget_core.power.time_domain import (
    BATTERY_DEPLETED,
    DOD_EXCEEDED,
    ORBIT_BALANCE_NEGATIVE,
    PEAK_POWER_EXCEEDED,
    time_domain_budget,
)
from tests.power_helpers import power_system, sv, synthetic_env, synthetic_run, td_project

PERIOD = 6000.0


def solve(env: EnvironmentData, *, system=None, rules=None, mode="a", **kw):  # type: ignore[no-untyped-def]
    run = synthetic_run(env, default_mode=mode, rules=rules)
    return time_domain_budget(td_project(system=system), run, **kw)


def with_battery(**changes: object):  # type: ignore[no-untyped-def]
    system = power_system()
    return system.model_copy(update={"battery": system.battery.model_copy(update=changes)})


def with_array(**changes: object):  # type: ignore[no-untyped-def]
    system = power_system()
    return system.model_copy(update={"solar_array": system.solar_array.model_copy(update=changes)})


# ---- PWR-1: BOL needs neither the degradation nor the capacity fade --------------------------


def test_bol_does_not_need_the_annual_degradation() -> None:
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    result = solve(env, system=with_array(annual_degradation_ratio=sv(None)))
    bol, eol = result.case("bol"), result.case("eol")
    assert bol.generation_w is not None and bol.energy_wh is not None
    assert bol.summary.generated_wh == pytest.approx(4200.0 * 100.0 / 3600.0)
    assert eol.generation_w is None and eol.energy_wh is None  # EOL does need it
    assert any(p.path == "solar_array.annual_degradation_ratio" for p in result.problems)


def test_bol_does_not_need_the_battery_capacity_fade() -> None:
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    result = solve(env, system=with_battery(annual_capacity_fade_ratio=sv(None)))
    bol, eol = result.case("bol"), result.case("eol")
    assert bol.capacity_wh == pytest.approx(100.0)
    assert bol.energy_wh is not None and bol.summary.maximum_dod_ratio is not None
    assert eol.generation_w is not None  # the array does not depend on the fade
    assert eol.capacity_wh is None and eol.energy_wh is None  # the pack does
    assert any(p.path == "battery.annual_capacity_fade_ratio" for p in result.problems)


# ---- PWR-2: battery saturation instants are refined edges ------------------------------------

# Mode b (100 W) in eclipse, mode a (40 W) in sunlight, eclipse 1000..4800 s. The array gives
# 100 W. The 100 Wh battery empties after 3600 s in the eclipse (t = 4600 s) and the sunlight
# recharges it at 60 W.
SAT_ECLIPSE = (1000.0, 4800.0)
SAT_RULES = [ScenarioRule(kind="in_eclipse", mode="b")]


@pytest.mark.parametrize("step", [10.0, 100.0, 1400.0, 4000.0, 8000.0])
def test_depletion_start_and_worst_value_do_not_depend_on_the_step(step: float) -> None:
    env = synthetic_env(2 * PERIOD, step, [SAT_ECLIPSE])
    bol = solve(env, rules=SAT_RULES).case("bol")
    (depleted,) = [v for v in bol.violations if v.code == BATTERY_DEPLETED]
    assert depleted.start_s == pytest.approx(4600.0, abs=1e-6)
    assert depleted.end_s == pytest.approx(4800.0, abs=1e-6)
    assert depleted.extreme == pytest.approx(100.0)
    assert bol.summary.unmet_wh == pytest.approx(100.0 * 200.0 / 3600.0)


@pytest.mark.parametrize("step", [10.0, 100.0, 1400.0, 4000.0, 8000.0])
def test_depth_of_discharge_crossing_is_exact_when_the_battery_empties_in_the_step(
    step: float,
) -> None:
    # limit 0.5: crossed 50 Wh / 100 W = 1800 s after the eclipse starts; a linear interpolation
    # across the step that empties the battery would say 2900 s at step 8000
    system = with_battery(max_dod_ratio={"eol": sv(0.5)})
    env = synthetic_env(2 * PERIOD, step, [SAT_ECLIPSE])
    bol = solve(env, system=system, rules=SAT_RULES).case("bol")
    (found,) = [v for v in bol.violations if v.code == DOD_EXCEEDED]
    assert found.start_s == pytest.approx(2800.0, abs=1e-6)
    assert found.extreme == pytest.approx(1.0)
    # sunlight at 60 W brings the stored energy back to 50 Wh 50 / 60 h = 3000 s after 4800 s
    assert found.end_s == pytest.approx(7800.0, abs=1e-6)


@pytest.mark.parametrize("step", [10.0, 700.0, 8000.0])
def test_full_battery_instant_does_not_change_the_energy_series(step: float) -> None:
    env = synthetic_env(2 * PERIOD, step, [SAT_ECLIPSE])
    bol = solve(env, rules=SAT_RULES).case("bol")
    # full and 60 W surplus for the first 1000 s; 100 W demand for 3600 s drains it; 60 W for
    # 6000 s refills it: full again at t = 10800 s, so 1000 + 1200 s of 60 W are curtailed
    assert bol.summary.curtailed_wh == pytest.approx(60.0 * 2200.0 / 3600.0)
    assert bol.summary.energy_change_wh == pytest.approx(0.0, abs=1e-6)


# ---- PWR-3 and PWR-11: orbit windows ---------------------------------------------------------


@pytest.mark.parametrize("step", [3000.0, 4000.0, 6000.0])
def test_no_orbit_windows_when_the_step_is_too_coarse(step: float) -> None:
    env = synthetic_env(4 * PERIOD, step, [(3000.0, 4800.0)])
    assert orbit_windows(env.grid.times_s, env.position_m) == []
    windows, reason = orbit_windows_with_reason(env.grid.times_s, env.position_m)
    assert windows == [] and reason in ("coarse", "plane")


def test_coarse_step_says_the_orbit_balance_is_not_evaluated() -> None:
    env = synthetic_env(4 * PERIOD, 4000.0, [(3000.0, 4800.0)])
    result = solve(env)
    assert result.windows == ()
    (problem,) = [p for p in result.problems if p.code == "ORBIT_BALANCE_NOT_EVALUATED"]
    assert problem.severity.value == "warning"
    assert "step" in problem.message
    assert {(n.case, n.code) for n in result.not_evaluated} >= {
        ("bol", ORBIT_BALANCE_NEGATIVE),
        ("eol", ORBIT_BALANCE_NEGATIVE),
    }


def test_a_step_of_a_quarter_orbit_still_gives_windows() -> None:
    env = synthetic_env(2 * PERIOD, 1500.0, [(3000.0, 4800.0)])
    assert len(orbit_windows(env.grid.times_s, env.position_m)) == 2


def test_no_complete_orbit_is_recorded_without_a_problem() -> None:
    env = synthetic_env(3000.0, 100.0, [])
    result = solve(env)
    assert not [p for p in result.problems if p.code == "ORBIT_BALANCE_NOT_EVALUATED"]
    (found,) = [n for n in result.not_evaluated if n.case == "bol"]
    assert found.code == ORBIT_BALANCE_NEGATIVE and "complete orbit" in found.reason


def test_the_last_orbit_is_kept_when_the_duration_is_not_a_multiple_of_the_step() -> None:
    # two orbits of 6000 s with a 7 s step: the last sample is at 11998 s, the scenario ends at
    # 12000 s; the second window must end at 12000 s
    env = synthetic_env(12000.0, 7.0, [(3000.0, 4800.0)])
    assert env.grid.times_s[-1] < 12000.0
    result = solve(env)
    assert len(result.windows) == 2
    assert result.windows[1][1] == pytest.approx(12000.0, abs=1.0)


# ---- PWR-4 and RC-7: checks that were not evaluated -------------------------------------------


def test_placeholder_limits_are_recorded_as_not_evaluated() -> None:
    system = with_battery(max_dod_ratio={"eol": sv(None)})
    system = system.model_copy(
        update={"limits": system.limits.model_copy(update={"peak_power_w": sv(None)})}
    )
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    result = solve(env, system=system)
    pairs = {(n.case, n.code) for n in result.not_evaluated}
    assert (None, PEAK_POWER_EXCEEDED) in pairs
    assert ("bol", DOD_EXCEEDED) in pairs and ("eol", DOD_EXCEEDED) in pairs
    assert ("bol", BATTERY_DEPLETED) not in pairs  # that check ran
    assert all(n.reason for n in result.not_evaluated)


def test_a_complete_run_has_nothing_that_was_not_evaluated() -> None:
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    assert solve(env).not_evaluated == ()


def test_missing_mission_phase_means_the_dod_limit_was_not_checked() -> None:
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    run = synthetic_run(env, phase=None)
    result = time_domain_budget(td_project(), run)  # two phases: none is chosen
    assert {n.code for n in result.not_evaluated if n.case == "bol"} == {DOD_EXCEEDED}


def test_violation_count_includes_the_case_independent_findings() -> None:
    rules = [ScenarioRule(kind="in_eclipse", mode="b")]  # 120 W peak against 110 W
    env = synthetic_env(PERIOD, 100.0, [(3000.0, 4800.0)])
    result = solve(env, rules=rules)
    peaks = [v for v in result.violations if v.code == PEAK_POWER_EXCEEDED]
    assert len(peaks) == 1 and peaks[0].case is None
    bol = result.case("bol")
    assert not [v for v in bol.violations if v.code == PEAK_POWER_EXCEEDED]
    assert result.violation_count("bol") == len(bol.violations) + 1


# ---- PWR-5: the soft shadow is the same at every step ------------------------------------------


def soft_env(step: float) -> EnvironmentData:
    """Eclipse 3000..4800 s with a 10 s penumbra at both ends (umbra 3010..4790 s)."""
    env = synthetic_env(PERIOD, step, [(3000.0, 4800.0)])
    t = env.grid.times_s
    ratio = np.ones_like(t)
    ratio[(t >= 3010.0) & (t < 4790.0)] = 0.0
    enter = (t >= 3000.0) & (t < 3010.0)
    ratio[enter] = 1.0 - (t[enter] - 3000.0) / 10.0
    leave = (t >= 4790.0) & (t < 4800.0)
    ratio[leave] = (t[leave] - 4790.0) / 10.0
    return replace(
        env,
        shadow_model="conical",
        sunlight_ratio=ratio,
        eclipses=(Interval(3000.0, 4800.0),),
        umbras=(Interval(3010.0, 4790.0),),
    )


@pytest.mark.parametrize("step", [1.0, 5.0, 10.0, 300.0, 1000.0, 2000.0])
def test_soft_shadow_energy_does_not_depend_on_the_step(step: float) -> None:
    bol = solve(soft_env(step)).case("bol")
    # sunlit 4200 s; each penumbra of 10 s counts half: 4210 s at 100 W
    assert bol.summary.generated_wh == pytest.approx(4210.0 * 100.0 / 3600.0, rel=1e-9)


def test_sharp_shadow_is_unchanged() -> None:
    bol = solve(synthetic_env(PERIOD, 1000.0, [(3000.0, 4800.0)])).case("bol")
    assert bol.summary.generated_wh == pytest.approx(4200.0 * 100.0 / 3600.0)


# ---- PWR-9: noise is not a depletion ------------------------------------------------------------


def test_generation_equal_to_demand_with_an_empty_battery_is_not_depleted() -> None:
    # 100 W in sunlight against mode b at 100 W; the battery starts empty. The step means of the
    # signals carry rounding noise far above 1e-9 W over a long run; it must not be a finding.
    system = with_battery(initial_soc_ratio=sv(0.0))
    env = synthetic_env(3 * 86400.0, 1.0, [])
    result = solve(env, system=system, mode="b")
    bol = result.case("bol")
    assert not [v for v in bol.violations if v.code == BATTERY_DEPLETED]


# ---- PWR-12 -------------------------------------------------------------------------------------


def test_unknown_case_name_is_refused() -> None:
    env = synthetic_env(PERIOD, 100.0, [])
    with pytest.raises(ValueError, match="case"):
        solve(env, cases=("bol", "mid"))


def test_violation_times_are_plain_floats() -> None:
    env = synthetic_env(2 * PERIOD, 100.0, [SAT_ECLIPSE])
    result = solve(env, rules=SAT_RULES)
    for v in result.violations:
        assert type(v.start_s) is float and type(v.end_s) is float


def test_missing_key_is_not_reported_as_a_missing_file() -> None:
    system = with_battery(max_dod_ratio={"launch": sv(0.5)})  # no entry for the phase "eol"
    env = synthetic_env(PERIOD, 100.0, [])
    result = solve(env, system=system)
    (problem,) = [p for p in result.problems if p.path == "battery.max_dod_ratio.eol"]
    assert "not provided" not in problem.message
    assert "config/power_system.yaml" in problem.message


# ---- PWR-4 and RC-7 in the command line and the report -----------------------------------------

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
FIXED = ["--user", "Test User", "--date", "2026-01-02T03:04:05Z"]


@pytest.fixture
def eps(tmp_path: Path) -> Path:
    shutil.copytree(EXAMPLES / "microsat_250kg_eps", tmp_path / "p")
    return tmp_path / "p"


def test_cli_counts_the_case_independent_violations_in_each_case(
    eps: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out"
    assert main(["power-timeline", str(eps), "--out", str(out), "--report", "json", *FIXED]) == 1
    text = capsys.readouterr().out
    data = json.loads((out / "power_time_one_day.json").read_text())
    peaks = [v for v in data["violations"] if v["case"] is None]
    assert peaks  # the example's peak power is above its limit
    bol = next(line for line in text.splitlines() if line.startswith("BOL:"))
    assert f"{len(peaks)} violation(s)" in bol  # no case finding of its own in the example


def test_cli_and_report_say_not_evaluated_instead_of_no_violations(
    eps: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = eps / "config" / "power_system.yaml"
    text = path.read_text()
    text = text.replace(
        "value: 180.0\n    source: Synthetic example value; not from a data sheet or a standard.",
        "value: null\n    source: TBD",
    )
    path.write_text(text)
    out = tmp_path / "out"
    main(["power-timeline", str(eps), "--out", str(out), "--report", "xlsx", *FIXED])
    shown = capsys.readouterr().out
    assert "Not evaluated:" in shown and "Peak power above the limit" in shown
    sheet = load_workbook(out / "power_time_one_day.xlsx")["Summary"]
    cells = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "Not evaluated:" in cells and "No violations found." not in cells


# ---- MI-16: bounded counts and directions ----------------------------------------------------


@pytest.mark.parametrize("field", ["strings", "cells_per_string"])
@pytest.mark.parametrize("count", [0, 10_001, 10**400])
def test_array_counts_are_bounded(field: str, count: int) -> None:
    values = {"name": "top", "normal_body": [0.0, 0.0, 1.0], "strings": 2, "cells_per_string": 2}
    with pytest.raises(ValidationError):
        ArrayFace(**{**values, field: count})  # type: ignore[arg-type]
    assert ArrayFace(**{**values, field: 10_000}).cell_count >= 10_000  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("field", "limit"), [("cells_in_series", 1_000), ("cells_in_parallel", 10_000)]
)
def test_battery_counts_are_bounded(field: str, limit: int) -> None:
    battery = power_system().battery
    with pytest.raises(ValidationError):
        Battery(**{**battery.model_dump(), field: limit + 1})
    assert getattr(Battery(**{**battery.model_dump(), field: limit}), field) == limit


@pytest.mark.parametrize(
    "vector", [[1e300, 0.0, 0.0], [0.0, 0.0, 1e7], [0.0, 0.0, 1e-9], [0, 0, 0]]
)
def test_directions_must_have_a_sane_length(vector: list[float]) -> None:
    with pytest.raises(ValidationError):
        ArrayFace(name="top", normal_body=vector, strings=1, cells_per_string=1)
    with pytest.raises(ValidationError):
        Attitude(default="sun", sun_direction_body=vector)


def test_face_cosines_survive_huge_and_tiny_vectors() -> None:
    from budget_core.power.array import face_cosines

    normals = np.array([[1e300, 0.0, 0.0], [0.0, 1e-300, 0.0]])
    sun = np.array([[1.0, 1.0, 0.0], [1e200, 1e200, 0.0]])
    assert face_cosines(normals, sun) == pytest.approx(math.sqrt(0.5))
