"""Unit power data as a table: export a template, fill it in a spreadsheet, import it back.

The table has one row per power mode of a unit. An import is planned first (`plan_import` lists
every change and every problem and writes nothing); `apply_plan` then edits the unit files, all of
them or none. Rules (decision D-102):

* an empty cell means "unchanged", never zero;
* only power modes of units that exist are touched; units are not created and modes are not
  deleted; a row for a new power mode of an existing unit adds it (average and peak are required);
* values may carry a unit ("800 mW", "40 %") and are converted like in the YAML files;
* a spreadsheet formula is never evaluated; a cell without a stored value counts as empty;
* messages name the file, the row (line) and the column, never a value read from a file.
"""

from __future__ import annotations

import csv
import io
import math
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from pydantic import ValidationError
from ruamel.yaml import YAML
from ruamel.yaml.comments import CommentedMap, CommentedSeq

from budget_core.io.errors import explain
from budget_core.io.paths import resolve_inside
from budget_core.io.yamlio import LineMap, YamlSyntaxError, load_yaml_text
from budget_core.model import PowerMode, Project
from budget_core.problems import Problem, Severity
from budget_core.reports.csvutil import csv_writer, defuse
from budget_core.reports.files import OutputError, write_file
from budget_core.units.quantity import UnitError, parse_quantity

KEY_COLUMNS = ("unit_id", "power_mode")
INFO_COLUMNS = ("unit_name", "subsystem")  # written for the reader, ignored on import
NUMBER_COLUMNS = (
    "avg_power_w",
    "peak_power_w",
    "duty_cycle_ratio",
    "heat_dissipation_ratio",
    "min_duration_s",
    "max_duration_s",
)
HEADER = ("unit_id", "unit_name", "subsystem", "power_mode", *NUMBER_COLUMNS)
SHEET = "Unit power modes"

MAX_FILE_BYTES = 5_000_000
MAX_UNPACKED_BYTES = 50_000_000  # an .xlsx is a zip; refuse one that unpacks to something huge
MAX_ROWS = 100_000
MAX_COLUMNS = 50
FIELD_ORDER = ("name", *NUMBER_COLUMNS)  # the order of the keys of a power mode in a unit file

Cell = str | float | None


# ---- export -----------------------------------------------------------------------------------


def _rows(project: Project) -> list[list[Any]]:
    out: list[list[Any]] = []
    for uid in sorted(project.units):
        unit = project.units[uid]
        for mode in unit.modes:
            out.append(
                [uid, unit.name, unit.subsystem, mode.name]
                + [getattr(mode, column) for column in NUMBER_COLUMNS]
            )
    return out


def table_csv(project: Project) -> str:
    """The template as CSV text (comma separated, LF). Text that a spreadsheet would run as a
    formula gets a leading apostrophe."""
    handle = io.StringIO()
    writer = csv_writer(handle)
    writer.writerow(HEADER)
    writer.writerows([["" if v is None else v for v in row] for row in _rows(project)])
    return handle.getvalue()


def table_xlsx(project: Project) -> bytes:
    """The template as an .xlsx workbook: one sheet of values (never a formula) and a sheet of
    instructions."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = SHEET
    sheet.append(list(HEADER))
    for row in _rows(project):
        sheet.append([defuse(v) if isinstance(v, str) else v for v in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(HEADER, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            14, len(name) + 2
        )
    notes = book.create_sheet("How to use")
    for line in (
        "One row per power mode of a unit. Fill in or change values, save, and import the file.",
        "An empty cell means: leave this value as it is. It does not mean zero.",
        "unit_id and power_mode identify the row. unit_name and subsystem are for reading only.",
        "A value may carry a unit, for example 800 mW or 40 %. Other columns are ignored.",
        "A row with a new power_mode adds that mode (average and peak power are needed).",
        "Units and power modes are never deleted by an import.",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ---- reading a table --------------------------------------------------------------------------


@dataclass(frozen=True)
class TableRow:
    line: int  # row number in the file (the header is row 1)
    cells: dict[str, Cell]


@dataclass(frozen=True)
class TableRead:
    rows: tuple[TableRow, ...]
    columns: tuple[str, ...]
    problems: tuple[Problem, ...]


def _error(code: str, message: str, file: str, hint: str = "", line: int | None = None) -> Problem:
    return Problem(Severity.ERROR, code, message, file=file, line=line, hint=hint)


def _unreadable(file: str, why: str = "") -> TableRead:
    message = "The file cannot be read as a table." + (f" {why}" if why else "")
    return TableRead(
        (),
        (),
        (
            _error(
                "IMPORT_FILE_UNREADABLE",
                message,
                file,
                "Use a .csv file (UTF-8) or an .xlsx workbook, for example one made by "
                "'budget units-export'.",
            ),
        ),
    )


def _too_large(file: str) -> TableRead:
    return TableRead(
        (),
        (),
        (
            _error(
                "IMPORT_FILE_TOO_LARGE",
                "The file is too large for a unit table.",
                file,
                "A unit table is a few kilobytes; check that this is the right file.",
            ),
        ),
    )


def _normal(cell: Any) -> Cell:
    if cell is None:
        return None
    if isinstance(cell, bool):
        return str(cell)  # refused later as a value; a yes/no is not a number
    if isinstance(cell, int | float):
        return float(cell)
    text = str(cell).strip()
    return text or None


def _delimiter(first_line: str) -> str:
    counts = {d: first_line.count(d) for d in (",", ";", "\t")}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] else ","


def _csv_rows(data: bytes) -> list[tuple[int, list[Cell]]] | None:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return None
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(text.split("\n", 1)[0]))
    rows: list[tuple[int, list[Cell]]] = []
    try:
        for record in reader:
            rows.append((reader.line_num, [_normal(c) for c in record]))
            if len(rows) > MAX_ROWS:
                raise OverflowError
    except csv.Error:
        return None
    return rows


def _xlsx_rows(path: Path) -> list[tuple[int, list[Cell]]] | None:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise OverflowError
        book = load_workbook(path, read_only=True, data_only=True)  # data_only: no formulas
    except OverflowError:
        raise
    except Exception:  # not a zip, not a workbook, damaged: all the same to the user
        return None
    rows: list[tuple[int, list[Cell]]] = []
    try:
        sheet = book.worksheets[0]
        for number, record in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append((number, [_normal(c) for c in record]))
            if len(rows) > MAX_ROWS:
                raise OverflowError
    except OverflowError:
        raise
    except Exception:
        return None
    finally:
        book.close()
    return rows


def read_table(path: Path) -> TableRead:
    """Read a .csv or .xlsx unit table. Nothing from the file ends up in a message."""
    path = Path(path)
    name = path.name
    suffix = path.suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        return _unreadable(name, "Only .csv and .xlsx files are read.")
    try:
        size = path.stat().st_size
    except OSError:
        return _unreadable(name, "The file does not exist or may not be read.")
    if size > MAX_FILE_BYTES:
        return _too_large(name)
    try:
        raw = _csv_rows(path.read_bytes()) if suffix == ".csv" else _xlsx_rows(path)
    except OverflowError:
        return _too_large(name)
    except OSError:
        return _unreadable(name, "The file does not exist or may not be read.")
    if raw is None:
        return _unreadable(name)
    return _structure(raw, name)


def _structure(raw: list[tuple[int, list[Cell]]], file: str) -> TableRead:
    if not raw:
        return TableRead((), (), (_error("IMPORT_FILE_UNREADABLE", "The table is empty.", file),))
    _, header_cells = raw[0]
    if len(header_cells) > MAX_COLUMNS:
        return _too_large(file)
    names = [str(c).strip().lower() if c is not None else "" for c in header_cells]
    problems: list[Problem] = []
    known = set(HEADER)
    for position, column in enumerate(names, start=1):
        if column and column not in known:
            problems.append(
                Problem(
                    Severity.WARNING,
                    "IMPORT_COLUMN_IGNORED",
                    f"Column {position} is not a column of the template and is ignored.",
                    file=file,
                    line=1,
                    hint="Columns of the template: " + ", ".join(HEADER) + ".",
                )
            )
    rows: list[TableRow] = []
    for line, cells in raw[1:]:
        values = {
            name: (cells[i] if i < len(cells) else None)
            for i, name in enumerate(names)
            if name in known
        }
        if any(v is not None for v in values.values()):
            rows.append(TableRow(line, values))
    return TableRead(tuple(rows), tuple(n for n in names if n in known), tuple(problems))


# ---- planning an import -----------------------------------------------------------------------


@dataclass(frozen=True)
class Change:
    unit_id: str
    power_mode: str
    field: str
    old: float | None  # None for a mode that is added
    new: float
    kind: Literal["changed", "added"]
    line: int  # row in the imported file


@dataclass(frozen=True)
class _Edit:
    unit_id: str
    power_mode: str
    values: dict[str, float]
    is_new: bool


@dataclass(frozen=True)
class ImportPlan:
    changes: tuple[Change, ...]
    unchanged: int  # power modes in the table that already have these values
    problems: tuple[Problem, ...]
    edits: tuple[_Edit, ...] = field(default=(), repr=False)

    @property
    def ok(self) -> bool:
        return not any(p.severity is Severity.ERROR for p in self.problems)


def _number(cell: Cell, column: str) -> float:
    if isinstance(cell, float):
        if not math.isfinite(cell):
            raise UnitError("the value must be a finite number")
        return cell
    assert cell is not None
    if "_" not in cell:  # Python reads "1_000" as a number; a spreadsheet does not
        try:
            plain = float(cell)  # a bare number is in the unit of the column, as in the YAML files
        except ValueError:
            return parse_quantity(cell, column)
        if not math.isfinite(plain):
            raise UnitError("the value must be a finite number")
        return plain
    return parse_quantity(cell, column)


def _same(a: float | None, b: float) -> bool:
    return a is not None and math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-15)


def plan_import(project: Project, table: TableRead, file: str) -> ImportPlan:
    """What importing `table` would change. Writes nothing."""
    problems: list[Problem] = list(table.problems)
    if any(p.severity is Severity.ERROR for p in problems):
        return ImportPlan((), 0, tuple(problems))
    for column in KEY_COLUMNS:
        if column not in table.columns:
            problems.append(
                _error(
                    "IMPORT_COLUMN_MISSING",
                    f"The table has no '{column}' column.",
                    file,
                    "Start from the template written by 'budget units-export'.",
                    line=1,
                )
            )
    if any(p.severity is Severity.ERROR for p in problems):
        return ImportPlan((), 0, tuple(problems))

    changes: list[Change] = []
    edits: list[_Edit] = []
    unchanged = 0
    seen: set[tuple[str, str]] = set()
    for row in table.rows:
        before = len(problems)
        unit_id = row.cells.get("unit_id")
        mode_name = row.cells.get("power_mode")
        if not isinstance(unit_id, str) or unit_id not in project.units:
            problems.append(
                Problem(
                    Severity.ERROR,
                    "IMPORT_UNKNOWN_UNIT",
                    "This row names a unit that the project does not have.",
                    file=file,
                    path="unit_id",
                    line=row.line,
                    hint="An import creates no units; use the ids of the exported template.",
                )
            )
            continue
        if not isinstance(mode_name, str):
            problems.append(
                Problem(
                    Severity.ERROR,
                    "IMPORT_FIELD_MISSING",
                    "This row has no power mode name.",
                    file=file,
                    path="power_mode",
                    line=row.line,
                    hint="Give the name of the power mode, as in the exported template.",
                )
            )
            continue
        if (unit_id, mode_name) in seen:
            problems.append(
                Problem(
                    Severity.ERROR,
                    "IMPORT_DUPLICATE_ROW",
                    "A row for this unit and power mode appears more than once.",
                    file=file,
                    path="power_mode",
                    line=row.line,
                    hint="Keep one row per power mode.",
                )
            )
            continue
        seen.add((unit_id, mode_name))

        given: dict[str, float] = {}
        for column in NUMBER_COLUMNS:
            cell = row.cells.get(column)
            if cell is None:
                continue
            try:
                given[column] = _number(cell, column)
            except UnitError as exc:
                problems.append(
                    Problem(
                        Severity.ERROR,
                        "IMPORT_VALUE_INVALID",
                        str(exc),
                        file=file,
                        path=column,
                        line=row.line,
                        hint="Write a number, or a number with a matching unit such as '2.5 W'.",
                    )
                )
        if len(problems) > before:
            continue

        existing = next((m for m in project.units[unit_id].modes if m.name == mode_name), None)
        if existing is None:
            missing = [c for c in ("avg_power_w", "peak_power_w") if c not in given]
            for column in missing:
                problems.append(
                    Problem(
                        Severity.ERROR,
                        "IMPORT_FIELD_MISSING",
                        "A new power mode needs average and peak power.",
                        file=file,
                        path=column,
                        line=row.line,
                        hint="Fill in both cells, or use the name of an existing power mode.",
                    )
                )
            if missing:
                continue
            merged: dict[str, Any] = {"name": mode_name, **given}
        else:
            merged = {"name": mode_name} | {
                c: getattr(existing, c) for c in NUMBER_COLUMNS if getattr(existing, c) is not None
            }
            merged.update(given)
        try:
            PowerMode.model_validate(merged)
        except ValidationError as exc:
            for found in explain(exc, PowerMode, file, LineMap()):
                problems.append(
                    Problem(
                        Severity.ERROR,
                        "IMPORT_VALUE_INVALID",
                        found.message,
                        file=file,
                        path=found.path,
                        line=row.line,
                        hint="Correct this value; the message says what the field requires.",
                    )
                )
            continue

        edit_values: dict[str, float] = {}
        for column, value in given.items():
            old = getattr(existing, column) if existing is not None else None
            if existing is not None and _same(old, value):
                continue
            edit_values[column] = value
            changes.append(
                Change(
                    unit_id,
                    mode_name,
                    column,
                    old,
                    value,
                    "added" if existing is None else "changed",
                    row.line,
                )
            )
        if existing is not None and not edit_values:
            unchanged += 1
        if edit_values:
            edits.append(_Edit(unit_id, mode_name, edit_values, existing is None))
    return ImportPlan(tuple(changes), unchanged, tuple(problems), tuple(edits))


# ---- applying a plan --------------------------------------------------------------------------


def _insert_in_order(mapping: CommentedMap, key: str, value: float) -> None:
    """Add `key` at its place in the canonical order of a power mode, or change it in place."""
    if key in mapping:
        mapping[key] = value
        return
    rank = FIELD_ORDER.index(key)
    position = 0
    for index, existing in enumerate(mapping.keys()):
        if existing in FIELD_ORDER and FIELD_ORDER.index(existing) < rank:
            position = index + 1
    mapping.insert(position, key, value)


def _edited_text(text: str, edits: Iterable[_Edit]) -> str:
    yaml = YAML(typ="rt")
    yaml.preserve_quotes = True
    yaml.default_flow_style = False
    yaml.indent(mapping=2, sequence=4, offset=2)
    yaml.width = 4096
    document = yaml.load(text)
    modes = document["modes"]
    for edit in edits:
        target: CommentedMap | None = None
        for candidate in modes:
            if candidate.get("name") == edit.power_mode:
                target = candidate
                break
        if target is None:
            if not edit.is_new:
                raise OutputError("A unit file no longer matches the plan; plan the import again.")
            target = CommentedMap()
            target["name"] = edit.power_mode
            assert isinstance(modes, CommentedSeq)
            modes.append(target)
        for key in sorted(edit.values, key=FIELD_ORDER.index):
            _insert_in_order(target, key, edit.values[key])
    buffer = io.StringIO()
    yaml.dump(document, buffer)
    return buffer.getvalue()


def apply_plan(root: Path, plan: ImportPlan) -> list[Path]:
    """Write the planned changes into the unit files, minimal edits that keep comments and order.
    A plan with an error writes nothing; otherwise every file is prepared before the first is
    written, so a problem with one file cannot leave half an import."""
    if not plan.ok or not plan.edits:
        return []
    root = Path(root)
    by_unit: dict[str, list[_Edit]] = {}
    for edit in plan.edits:
        by_unit.setdefault(edit.unit_id, []).append(edit)
    prepared: list[tuple[Path, str]] = []
    for unit_id in sorted(by_unit):
        target = resolve_inside(root, f"units/{unit_id}.yaml")
        if target is None or not target.is_file():
            raise OutputError("A unit file could not be found; plan the import again.")
        try:
            text = target.read_text(encoding="utf-8")
            load_yaml_text(text)  # the same size and alias checks as the loader
            new_text = _edited_text(text, by_unit[unit_id])
            load_yaml_text(new_text)
        except (OSError, UnicodeDecodeError, YamlSyntaxError, KeyError, TypeError) as exc:
            raise OutputError("A unit file could not be edited; it was not changed.") from exc
        prepared.append((target, new_text))
    written: list[Path] = []
    for target, new_text in prepared:
        written.append(write_file(target.parent, target.name, new_text, sanitise=False))
    return written
