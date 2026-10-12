"""XLSX renderer (openpyxl). Output is byte-identical for identical documents."""

from __future__ import annotations

import io
import math
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from budget_core.reports.document import (
    NA,
    Column,
    ReportDocument,
    Section,
    Table,
    banner_headline,
    format_cell,
)
from budget_core.reports.numfmt import no_negative_zero
from budget_core.reports.textsafe import clean_document
from budget_core.reports.zipnorm import normalise_zip

FONT = "IBM Plex Sans"
HEADER_FILL = PatternFill("solid", fgColor="E0E0E0")
FIGURE_WIDTH_PX = 1100
ROW_PX = 20  # default row height in pixels
CELL_LIMIT = 32000  # a spreadsheet cell holds 32,767 characters


def _number_format(column: Column) -> str:
    if column.decimals < 0:
        return "General"
    if column.kind == "percent":
        return "0." + "0" * max(column.decimals - 2, 0) + "%" if column.decimals > 2 else "0%"
    return "0." + "0" * column.decimals if column.decimals else "0"


def _write_table(ws: Worksheet, table: Table, row: int, widths: dict[int, int]) -> int:
    ws.cell(row=row, column=1, value=table.title).font = Font(name=FONT, bold=True, size=12)
    row += 1
    for c, column in enumerate(table.columns, 1):
        cell = ws.cell(row=row, column=c, value=column.header)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths[c] = max(widths.get(c, 0), min(len(column.header), 40))
    ws.freeze_panes = ws.freeze_panes or None
    row += 1
    last = len(table.rows) - 1
    shortened = 0
    for i, values in enumerate(table.rows):
        for c, (value, column) in enumerate(zip(values, table.columns, strict=True), 1):
            if value is None or (not isinstance(value, str) and not math.isfinite(value)):
                cell = ws.cell(row=row, column=c, value=NA)
                cell.alignment = Alignment(horizontal="right")
                text = NA
            elif isinstance(value, str):
                if len(value) > CELL_LIMIT:
                    value = value[: CELL_LIMIT - 1] + "…"
                    shortened += 1
                cell = ws.cell(row=row, column=c, value=value)
                text = value
            else:
                cell = ws.cell(row=row, column=c, value=no_negative_zero(float(value)))
                cell.number_format = _number_format(column)
                text = format_cell(value, column)
            cell.font = Font(name=FONT, bold=table.emphasise_last_row and i == last)
            widths[c] = max(widths.get(c, 0), min(len(text), 60))
        row += 1
    if table.note:
        ws.cell(row=row, column=1, value=table.note).font = Font(name=FONT, italic=True)
        row += 1
    if shortened:
        note = (
            f"{shortened} cell(s) were shortened to fit a spreadsheet cell (32,767 characters); "
            "the DOCX and JSON files hold the full text."
        )
        ws.cell(row=row, column=1, value=note).font = Font(name=FONT, italic=True)
        row += 1
    return row + 1


def _write_section(ws: Worksheet, section: Section, banner: str, first: bool, title: str) -> None:
    row = 1
    if first:
        ws.cell(row=row, column=1, value=title).font = Font(name=FONT, bold=True, size=14)
        row += 1
    if banner:  # on every sheet: a sheet can be read or copied without the first one
        cell = ws.cell(row=row, column=1, value=banner if first else banner_headline(banner))
        cell.font = Font(name=FONT, bold=True, color="DA1E28")
        row += 1
    if first or banner:
        row += 1
    for text in section.paragraphs:
        ws.cell(row=row, column=1, value=text).font = Font(name=FONT)
        row += 1
    widths: dict[int, int] = {}
    for table in section.tables:
        row = _write_table(ws, table, row, widths)
    for figure in section.figures:
        ws.cell(row=row, column=1, value=figure.title).font = Font(name=FONT, bold=True, size=12)
        row += 1
        image = SheetImage(io.BytesIO(figure.png))
        scale = FIGURE_WIDTH_PX / image.width
        image.width, image.height = FIGURE_WIDTH_PX, round(image.height * scale)
        ws.add_image(image, f"A{row}")
        row += image.height // ROW_PX + 2
        if figure.alt:
            ws.cell(row=row, column=1, value=figure.alt).font = Font(name=FONT, italic=True)
            row += 2
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = max(width + 2, 10)


def render_xlsx(doc: ReportDocument) -> bytes:
    doc = clean_document(doc)
    wb = Workbook()
    wb.remove(wb.active)
    for i, section in enumerate(doc.sections):
        ws = wb.create_sheet(section.sheet_name)
        _write_section(ws, section, doc.banner, i == 0, doc.title)
    for ws in wb.worksheets:  # project text such as "=HYPERLINK(...)" must stay text
        for sheet_row in ws.iter_rows():
            for cell in sheet_row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    stamp = datetime.strptime(doc.provenance.generated, "%Y-%m-%dT%H:%M:%SZ")
    props = wb.properties
    props.creator = doc.provenance.user  # the author, as in the PDF and DOCX properties
    props.lastModifiedBy = doc.provenance.tool
    props.title = doc.title
    props.created = stamp
    props.modified = stamp
    buffer = io.BytesIO()
    # Not wb.save(): openpyxl's save_workbook overwrites `modified` with the current time.
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        ExcelWriter(wb, archive).write_data()
    return normalise_zip(buffer.getvalue())
